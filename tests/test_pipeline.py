import tempfile
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

from cm_power_merge.columns import to_number
from cm_power_merge.compare import compare_frames, load_output_rows
from cm_power_merge.pipeline import run_pipeline
from cm_power_merge.report import write_exceptions, write_workbook
from cm_power_merge.validation import overall_status, readiness_checks, run_validations
from helpers import MONTH, make_cm, make_pm, make_ref, make_config


def _run(cfg=None, pm=None, cm=None):
    cfg = cfg or make_config()
    res = run_pipeline(pm if pm is not None else make_pm(), cm if cm is not None else make_cm(),
                       cfg, MONTH, make_ref)
    return cfg, res


def test_merge_counts_and_overlay():
    cfg, res = _run()
    d = res.data.set_index("Service Asset Number")
    assert res.stats["kept_from_pm"] == 5 and res.stats["added_from_cm"] == 2 and res.stats["churned"] == 1
    assert res.churned["_SAN_"].tolist() == ["S5"]
    assert set(d.index) == {"S1", "S2", "S3", "S4", "S6", "S7", "S8"}
    assert d.loc["S1", "Gross Power MRR"] == 1100 and d.loc["S1", "PM MRR"] == 1000
    assert d.loc["S1", "PM MRR Var"] == 100 and abs(d.loc["S1", "COLA/Price Increase Check"] - 0.1) < 1e-12
    assert d.loc["S7", "PM MRR"] == 0 and d.loc["S7", "PM MRR Var"] == 700      # new SAN: variance = full value
    assert d.loc["S3", "Adjusted kW"] == 6                                        # recalculated from new Bill Qty
    assert d.loc["S2", "Adjusted kW"] == 0                                        # stale 99 kW on exempt row corrected
    assert d.loc["S4", "Adjusted kW"] == 0.44
    assert d.loc["S7", "Adjusted kW"] == 10                                       # Circuit kW 5 x Bill Qty 2
    assert d.loc["S7", "Commercial Model Helper"] == "Circuit" and d.loc["S8", "Commercial Model Helper"] == "Space"
    assert d.loc["S7", "Commercial Model"] in {"Circuit", "Per-kW", "+E"}
    assert d.loc["S1", "Commercial Model"] == "Circuit" and d.loc["S3", "Commercial Model"] == "+E"   # retained
    assert d.loc["S8", "Cabinet Count"] == 1 and d.loc["S1", "Customer Has Cabinet"] == "Yes"
    assert d.loc["S3", "Customer Has Cabinet"] == "No"
    assert res.data.columns[0] == "Reporting Month"
    assert "_SOURCE_" not in res.data.columns and "_SAN_" in res.data.columns
    assert len(res.kw_diffs) >= 2                                                  # S2 (99 -> 0) and S3 (5 -> 6)


def test_customers_ordered_by_mrr_then_key():
    _, res = _run()
    order = res.data["Customer Number"].astype(str).drop_duplicates().tolist()
    assert order == ["100", "300", "200", "400"][:1] + order[1:]                  # Acme (1450) first
    sums = res.data.groupby("Customer Number")["Monthly Amount"].sum()
    firsts = [sums[int(c)] if int(c) in sums.index else sums[c] for c in order]
    assert firsts == sorted(firsts, reverse=True)


def test_validation_passes_and_flags_expected_items():
    cfg, res = _run()
    checks = run_validations(res, cfg)
    by = {c.name: c for c in checks}
    assert by["completeness"].status == "PASS" and by["mrr_reconciles_to_cm"].status == "PASS"
    assert by["commercial_model_values"].status == "PASS"
    assert by["kw_zero_when_mrr_zero"].status == "PASS" and by["one_u_kw"].status == "PASS"
    assert by["kw_populated_for_revenue_rows"].status == "PASS"
    assert overall_status(checks) in ("PASS", "WARN")
    assert not any(c.status == "FAIL" for c in checks), [c for c in checks if c.status == "FAIL"]


def test_deployment_flag_is_na_without_pkw():
    _, res = _run()
    assert res.data["Per kW >60%±10% OK"].isna().all()                           # no PKW anywhere -> N/A, not False


def test_duplicate_cm_san_uses_last_row_and_warns():
    cm = make_cm()
    dup = cm[cm["Service Asset Number"] == "S3"].copy()
    dup["Monthly Amount"], dup["Gross Power MRR"] = 650, 650
    cfg, res = _run(cm=pd.concat([cm, dup], ignore_index=True))
    assert res.data.set_index("Service Asset Number").loc["S3", "Monthly Amount"] == 650
    by = {c.name: c for c in run_validations(res, cfg)}
    assert by["cm_san_unique"].status == "WARN" and by["completeness"].status == "PASS"
    assert by["mrr_reconciles_to_cm"].status == "PASS"                            # dedup before reconciling
    cfg.rules.duplicate_san_policy = "fail"
    assert {c.name: c for c in run_validations(res, cfg)}["cm_san_unique"].status == "FAIL"


def test_missing_cm_column_is_reported_not_guessed():
    cm = make_cm().drop(columns=["Monthly Amount"])
    cfg, res = _run(cm=cm)
    assert "Monthly Amount" in res.stats["missing_cm_targets"]
    by = {c.name: c for c in run_validations(res, cfg)}
    assert by["required_cm_columns"].status == "FAIL"


def test_kw_policy_preserve_existing_matches_legacy():
    cfg = make_config()
    cfg.rules.kw_policy = "preserve_existing"
    _, res = _run(cfg)
    d = res.data.set_index("Service Asset Number")
    assert d.loc["S2", "Adjusted kW"] == 99.0 and d.loc["S2", "Adjusted kW Source"] == "existing"
    assert d.loc["S3", "Adjusted kW"] == 5                                         # stale value kept, but listed for review
    assert "S3" in res.kw_diffs["_SAN_"].tolist()


def test_readiness_gate():
    cfg = make_config()
    assert readiness_checks(0, 100, cfg)[0].status == "FAIL"
    assert any(c.status == "FAIL" for c in readiness_checks(50, 100, cfg))
    assert all(c.status == "PASS" for c in readiness_checks(102, 100, cfg))


def test_workbook_has_formulas_with_cached_values_and_roundtrips():
    cfg, res = _run()
    checks = run_validations(res, cfg)
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "out.xlsx"
        write_workbook(path, res, checks, {"Run ID": "t"}, [("WARNING", "x")], cfg)
        assert path.exists() and not (Path(tmp) / "out.xlsx.tmp").exists()
        cols = list(res.data.columns)
        ws_f = openpyxl.load_workbook(path)["Merged"]
        ws_v = openpyxl.load_workbook(path, data_only=True)["Merged"]
        gp = cols.index("Gross Power MRR") + 1
        sub_rows = [r for r in range(2, ws_v.max_row + 1)
                    if ws_v.cell(r, cols.index("Customer Name") + 1).value == "zz Total"]
        assert len(sub_rows) == res.data["Customer Number"].nunique()
        first_sub = sub_rows[0]
        formula = ws_f.cell(first_sub, gp).value
        assert formula.startswith("=_xlfn.AGGREGATE(9,6,") and f"2:" in formula
        cust = res.data["Customer Number"].iloc[0]
        expected = res.data.loc[res.data["Customer Number"] == cust, "Gross Power MRR"].sum()
        assert ws_v.cell(first_sub, gp).value == expected                          # cached value, not 0
        grand = max(r for r in range(2, ws_v.max_row + 1) if ws_v.cell(r, 2).value == "Total")
        assert ws_v.cell(grand, gp).value == res.data["Gross Power MRR"].sum()
        assert ws_f.cell(grand, gp).value.startswith("=_xlfn.AGGREGATE(9,3,")
        cola = cols.index("COLA/Price Increase Check") + 1
        assert ws_f.cell(2, cola).value.startswith("=IFERROR(")
        assert "Run Info" in openpyxl.load_workbook(path).sheetnames

        # the written workbook is a valid next-month PM input and a valid compare input
        again = load_output_rows(path)
        assert len(again) == len(res.data)
        result = compare_frames(again, res.data.assign(**{"Reporting Month": "x"}))
        assert result.summary["rows_only_in_new"] == 0 and result.summary["rows_only_in_reference"] == 0
        out_csv = write_exceptions(Path(tmp) / "exc", res)
        assert any(p.name == "churned_sans.csv" for p in out_csv)


def test_next_month_run_accepts_previous_output_as_pm():
    cfg, res = _run()
    checks = run_validations(res, cfg)
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "out.xlsx"
        write_workbook(path, res, checks, {}, [], cfg)
        pm_next = pd.read_excel(path, dtype=object)
        cfg2, res2 = _run(pm=pm_next, cm=make_cm())
        assert res2.stats["pm_rows"] == len(res.data) and res2.stats["added_from_cm"] == 0
        d = res2.data.set_index("Service Asset Number")
        assert d.loc["S1", "PM MRR"] == 1100 and d.loc["S1", "PM MRR Var"] == 0   # this month's figure is next month's PM


def test_compare_detects_changes():
    _, res = _run()
    ref = res.data.copy()
    changed = res.data.copy()
    changed.loc[changed["Service Asset Number"] == "S1", "Adjusted kW"] = 11
    changed = changed[changed["Service Asset Number"] != "S8"]
    out = compare_frames(changed, ref)
    assert out.summary["rows_only_in_reference"] == 1 and out.summary["columns_with_differences"]["Adjusted kW"] == 1
    assert compare_frames(ref, ref).summary["cells_different"] == 0


def _missing_ref():
    raise FileNotFoundError("Commercial Model Reference not found: x.xlsx")


def test_missing_model_ref_stops_run_by_default_but_can_be_optional():
    cfg = make_config()
    try:
        run_pipeline(make_pm(), make_cm(), cfg, MONTH, _missing_ref)
    except FileNotFoundError:
        pass
    else:
        raise AssertionError("required reference should stop the run")
    cfg.rules.model_ref_required = False
    res = run_pipeline(make_pm(), make_cm(), cfg, MONTH, _missing_ref)
    assert res.stats["new_rows_without_cmh"] == 2
    by = {c.name: c for c in run_validations(res, cfg)}
    assert by["new_rows_have_helper"].status == "FAIL"          # still visible, never silent


def test_reference_not_loaded_when_new_rows_already_have_a_helper():
    cm = make_cm()
    cm["Commercial Model Helper"] = cm["Base SKU"].map({"PWR-CIRC": "Circuit", "CAB-42U": "Space"})
    res = run_pipeline(make_pm(), cm, make_config(), MONTH, _missing_ref)     # would raise if it were loaded
    assert res.stats["new_rows_without_cmh"] == 0


def test_header_text_in_helper_column_counts_as_blank_and_loads_reference():
    cm = make_cm()
    cm["Commercial Model Helper"] = "Commercial Model Helper"      # header text leaked into the data
    res = run_pipeline(make_pm(), cm, make_config(), MONTH, make_ref)
    d = res.data.set_index("Service Asset Number")
    assert d.loc["S7", "Commercial Model Helper"] == "Circuit" and res.stats["new_rows_without_cmh"] == 0


def test_workbook_colours_and_number_formats():
    cfg = make_config()
    cfg.excel.number_columns = ["Monthly Amount", "Gross Power MRR"]
    cfg.excel.number_format = "#,##0.00"
    cfg.excel.header_bg, cfg.excel.subtotal_bg, cfg.excel.grand_total_bg = "#112233", "#AABBCC", "#CCDDEE"
    _, res = _run(cfg)
    checks = run_validations(res, cfg)
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "styled.xlsx"
        write_workbook(path, res, checks, {}, [], cfg)
        wb = openpyxl.load_workbook(path)
        ws, cols = wb["Merged"], list(res.data.columns)
        ma = cols.index("Monthly Amount") + 1
        assert ws.cell(1, 1).fill.fgColor.rgb.endswith("112233") and ws.cell(1, 1).font.bold
        sub = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, cols.index("Customer Name") + 1).value == "zz Total")
        assert ws.cell(sub, 1).fill.fgColor.rgb.endswith("AABBCC")              # blank cells are coloured too
        assert ws.cell(sub, ma).fill.fgColor.rgb.endswith("AABBCC") and ws.cell(sub, ma).number_format == "#,##0.00"
        assert ws.cell(2, ma).number_format == "#,##0.00"                          # data rows
        grand = max(r for r in range(2, ws.max_row + 1) if ws.cell(r, 2).value == "Total")
        assert ws.cell(grand, ma).fill.fgColor.rgb.endswith("CCDDEE") and ws.cell(grand, ma).number_format == "#,##0.00"
        assert ws.cell(2, cols.index("Customer Number") + 1).number_format == "General"   # untouched columns
        # cached totals still intact
        v = openpyxl.load_workbook(path, data_only=True)["Merged"]
        assert v.cell(grand, ma).value == res.data["Monthly Amount"].sum()
